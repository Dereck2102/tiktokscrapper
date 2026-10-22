"""Data serialization and export services for JSON, CSV, and Excel formats."""

from __future__ import annotations

import csv
import io
import json
from abc import ABC, abstractmethod
from typing import Any, Optional, Sequence

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    XLSX_DISPONIBLE = True
except ImportError:  # pragma: no cover
    XLSX_DISPONIBLE = False

COLUMNAS_POSTS: list[tuple[str, str]] = [
    ("id", "ID"),
    ("fecha", "Fecha"),
    ("dia_semana", "Día"),
    ("hora", "Hora"),
    ("tipo", "Tipo"),
    ("duracion", "Duración (s)"),
    ("fijado", "Fijado"),
    ("vistas", "Vistas"),
    ("likes", "Likes"),
    ("comentarios", "Comentarios"),
    ("compartidos", "Compartidos"),
    ("guardados", "Guardados"),
    ("interacciones", "Interacciones"),
    ("er", "ER %"),
    ("er_seguidores", "ER seguidores %"),
    ("vistas_por_dia", "Vistas/día"),
    ("hashtags", "Hashtags"),
    ("menciones", "Menciones"),
    ("musica", "Música"),
    ("musica_original", "Audio original"),
    ("descripcion", "Descripción"),
    ("url", "URL"),
    ("fuente", "Fuente"),
]


class PostRowNormalizer:
    """Flattens nested post dictionaries into tabular row dictionaries."""

    @staticmethod
    def to_rows(posts: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for post in posts:
            stats = post.get("estadisticas", {})
            music = post.get("musica") or {}
            rows.append(
                {
                    "id": post.get("id"),
                    "fecha": post.get("fecha"),
                    "dia_semana": post.get("dia_semana"),
                    "hora": post.get("hora"),
                    "tipo": post.get("tipo"),
                    "duracion": post.get("duracion"),
                    "fijado": "Sí" if post.get("fijado") else "No",
                    "vistas": stats.get("vistas"),
                    "likes": stats.get("likes"),
                    "comentarios": stats.get("comentarios"),
                    "compartidos": stats.get("compartidos"),
                    "guardados": stats.get("guardados"),
                    "interacciones": post.get("interacciones"),
                    "er": post.get("er"),
                    "er_seguidores": post.get("er_seguidores"),
                    "vistas_por_dia": post.get("vistas_por_dia"),
                    "hashtags": ", ".join(post.get("hashtags") or []),
                    "menciones": ", ".join(post.get("menciones") or []),
                    "musica": music.get("titulo"),
                    "musica_original": (
                        None
                        if music.get("original") is None
                        else ("Sí" if music.get("original") else "No")
                    ),
                    "descripcion": post.get("descripcion"),
                    "url": post.get("url"),
                    "fuente": post.get("fuente"),
                }
            )
        return rows


class ScrapeReportExporter(ABC):
    """Contract for exporting an individual scrape execution and its dashboard metrics."""

    @property
    @abstractmethod
    def mime_type(self) -> str:
        """Returns the HTTP MIME type for the exported artifact."""

    @abstractmethod
    def export_scrape(
        self, raw_result: dict[str, Any], dashboard: dict[str, Any]
    ) -> bytes:
        """Serializes a single scrape result and its computed dashboard."""


class GlobalDatasetExporter(ABC):
    """Contract for exporting the consolidated multi-account dataset."""

    @property
    @abstractmethod
    def mime_type(self) -> str:
        """Returns the HTTP MIME type for the exported artifact."""

    @abstractmethod
    def export_dataset(
        self, posts: list[dict[str, Any]], profiles: list[dict[str, Any]]
    ) -> bytes:
        """Serializes the global posts and profile history datasets."""


class JsonExporter(ScrapeReportExporter, GlobalDatasetExporter):
    """Serializes scrape reports and datasets as UTF-8 formatted JSON."""

    @property
    def mime_type(self) -> str:
        return "application/json"

    @staticmethod
    def serialize(data: Any) -> bytes:
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")

    def export_scrape(
        self, raw_result: dict[str, Any], dashboard: dict[str, Any]
    ) -> bytes:
        return self.serialize(raw_result)

    def export_dataset(
        self, posts: list[dict[str, Any]], profiles: list[dict[str, Any]]
    ) -> bytes:
        return self.serialize({"posts": posts, "perfiles": profiles})


class CsvExporter(ScrapeReportExporter, GlobalDatasetExporter):
    """Serializes tabular records as UTF-8 with BOM CSV streams."""

    @property
    def mime_type(self) -> str:
        return "text/csv"

    @staticmethod
    def serialize(
        rows: Sequence[dict[str, Any]],
        columns: Optional[Sequence[tuple[str, str]]] = None,
        delimiter: str = ",",
    ) -> bytes:
        if columns is None:
            columns = [(k, k) for k in (rows[0].keys() if rows else [])]
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=delimiter)
        writer.writerow([title for _, title in columns])
        for row in rows:
            writer.writerow(
                ["" if row.get(key) is None else row.get(key) for key, _ in columns]
            )
        return buffer.getvalue().encode("utf-8-sig")

    def export_scrape(
        self, raw_result: dict[str, Any], dashboard: dict[str, Any]
    ) -> bytes:
        rows = PostRowNormalizer.to_rows(dashboard["posts"])
        return self.serialize(rows, COLUMNAS_POSTS)

    def export_dataset(
        self, posts: list[dict[str, Any]], profiles: list[dict[str, Any]]
    ) -> bytes:
        return self.serialize(posts)


class ExcelExporter(ScrapeReportExporter, GlobalDatasetExporter):
    """Serializes multi-sheet OpenXML workbooks (.xlsx)."""

    @property
    def mime_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @staticmethod
    def _populate_sheet(
        workbook: Workbook,
        title: str,
        columns: Sequence[tuple[str, str]],
        rows: Sequence[dict[str, Any]],
        is_first: bool = False,
    ) -> None:
        sheet = workbook.active if is_first else workbook.create_sheet()
        sheet.title = title[:31]
        sheet.append([header for _, header in columns])
        for row in rows:
            sheet.append([row.get(key) for key, _ in columns])

        header_fill = PatternFill("solid", fgColor="FE2C55")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for idx, (key, header) in enumerate(columns, 1):
            max_len = max(
                [len(str(header))]
                + [len(str(r.get(key) or "")) for r in rows[:200]]
            )
            sheet.column_dimensions[get_column_letter(idx)].width = min(
                max(10, max_len + 2), 60
            )
        sheet.freeze_panes = "A2"
        if rows:
            sheet.auto_filter.ref = sheet.dimensions

    def export_scrape(
        self, raw_result: dict[str, Any], dashboard: dict[str, Any]
    ) -> bytes:
        if not XLSX_DISPONIBLE:
            raise RuntimeError("openpyxl no está instalado.")
        workbook = Workbook()
        profile = dashboard.get("perfil") or {}
        summary = [
            {"m": "Usuario", "v": profile.get("usuario")},
            {"m": "Nombre", "v": profile.get("nombre")},
            {"m": "Seguidores", "v": profile.get("seguidores")},
            {"m": "Siguiendo", "v": profile.get("siguiendo")},
            {"m": "Likes totales", "v": profile.get("likes_totales")},
            {"m": "Videos", "v": profile.get("videos")},
            {
                "m": "Fecha de scraping",
                "v": dashboard.get("meta", {}).get("fecha_scrape"),
            },
        ]
        summary.extend(
            {"m": key.replace("_", " ").capitalize(), "v": val}
            for key, val in dashboard["kpis"].items()
        )
        self._populate_sheet(
            workbook, "Resumen", [("m", "Métrica"), ("v", "Valor")], summary, is_first=True
        )
        self._populate_sheet(
            workbook,
            "Posts",
            COLUMNAS_POSTS,
            PostRowNormalizer.to_rows(dashboard["posts"]),
        )
        self._populate_sheet(
            workbook,
            "Insights",
            [("tipo", "Tipo"), ("titulo", "Insight"), ("texto", "Detalle")],
            dashboard["insights"],
        )
        self._populate_sheet(
            workbook,
            "Hashtags",
            [
                ("tag", "Hashtag"),
                ("usos", "Usos"),
                ("vistas_promedio", "Vistas prom."),
                ("er_promedio", "ER prom. %"),
                ("indice", "Índice vs mediana"),
            ],
            dashboard["hashtags"],
        )
        self._populate_sheet(
            workbook,
            "Por día",
            [("dia", "Día"), ("posts", "Posts"), ("vistas_promedio", "Vistas prom.")],
            dashboard["dias"],
        )
        self._populate_sheet(
            workbook,
            "Por hora",
            [("hora", "Hora"), ("posts", "Posts"), ("vistas_promedio", "Vistas prom.")],
            dashboard["horas"],
        )
        self._populate_sheet(
            workbook,
            "Duración",
            [
                ("rango", "Rango"),
                ("posts", "Posts"),
                ("vistas_promedio", "Vistas prom."),
                ("er_promedio", "ER prom. %"),
            ],
            dashboard["duraciones"],
        )
        if dashboard.get("historial"):
            self._populate_sheet(
                workbook,
                "Histórico",
                [
                    ("fecha", "Fecha"),
                    ("seguidores", "Seguidores"),
                    ("siguiendo", "Siguiendo"),
                    ("likes_totales", "Likes"),
                    ("videos", "Videos"),
                ],
                dashboard["historial"],
            )
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def export_dataset(
        self, posts: list[dict[str, Any]], profiles: list[dict[str, Any]]
    ) -> bytes:
        if not XLSX_DISPONIBLE:
            raise RuntimeError("openpyxl no está instalado.")
        workbook = Workbook()
        post_cols = [(k, k) for k in (posts[0].keys() if posts else ["id"])]
        prof_cols = [(k, k) for k in (profiles[0].keys() if profiles else ["username"])]
        self._populate_sheet(workbook, "posts", post_cols, posts, is_first=True)
        self._populate_sheet(workbook, "perfiles_historico", prof_cols, profiles)
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()


class ExporterRegistry:
    """Registry that resolves exporters by file format identifier."""

    def __init__(self) -> None:
        self._scrape_exporters: dict[str, ScrapeReportExporter] = {
            "json": JsonExporter(),
            "csv": CsvExporter(),
            "xlsx": ExcelExporter(),
        }
        self._dataset_exporters: dict[str, GlobalDatasetExporter] = {
            "json": JsonExporter(),
            "csv": CsvExporter(),
            "xlsx": ExcelExporter(),
        }

    def get_scrape_exporter(self, fmt: str) -> Optional[ScrapeReportExporter]:
        return self._scrape_exporters.get(fmt.lower())

    def get_dataset_exporter(self, fmt: str) -> Optional[GlobalDatasetExporter]:
        return self._dataset_exporters.get(fmt.lower())


default_exporter_registry = ExporterRegistry()


def posts_a_filas(posts: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    """Normalizes post dictionaries into flat tabular rows."""
    return PostRowNormalizer.to_rows(posts)


def a_json(datos: Any) -> bytes:
    """Serializes arbitrary data to UTF-8 JSON bytes."""
    return JsonExporter.serialize(datos)


def a_csv(
    filas: Sequence[dict[str, Any]],
    columnas: Optional[Sequence[tuple[str, str]]] = None,
    separador: str = ",",
) -> bytes:
    """Serializes row dictionaries to UTF-8-SIG CSV bytes."""
    return CsvExporter.serialize(filas, columnas, separador)


def dashboard_a_xlsx(dash: dict[str, Any]) -> bytes:
    """Serializes a dashboard model into a multi-sheet Excel workbook."""
    return ExcelExporter().export_scrape({}, dash)


def dataset_a_xlsx(
    posts: list[dict[str, Any]], perfiles: list[dict[str, Any]]
) -> bytes:
    """Serializes global posts and profile history into an Excel workbook."""
    return ExcelExporter().export_dataset(posts, perfiles)

